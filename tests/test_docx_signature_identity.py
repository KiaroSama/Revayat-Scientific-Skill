"""Package signature metadata is identified by OPC types, not only folder names."""
from pathlib import Path
import sys
import tempfile
import unittest.mock
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills/revayat-scientific/scripts'))
import docx_package as PACKAGE
from runtime import operation_log

PREFIX = 'application/vnd.openxmlformats-package.digital-signature-'
RELATIONSHIP = PACKAGE.R + '/digital-signature/'


class DocxSignatureIdentityTest(unittest.TestCase):
    def setUp(self):
        scratch = ROOT / '.scratch'
        scratch.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=scratch, prefix='docx signature metadata ')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source, self.target = self.root / 'source.docx', self.root / 'output.docx'
        scope = operation_log('test-docx-signature-identity', self.root / 'logs')
        self.log = scope.__enter__()
        self.addCleanup(scope.__exit__, None, None, None)
        self.log.info('running test=%s; metadata fixtures are not authenticated signatures', self._testMethodName)

    def fixture(self, *, typed=True, relationships=True, folder='signing', uppercase=False):
        from docx import Document
        document = Document()
        document.add_paragraph('Original text')
        document.save(self.source)
        with zipfile.ZipFile(self.source) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        origin, signature = folder + '/origin.sigs', folder + '/signature.data'
        members[origin] = b''
        members[signature] = b'<Signature xmlns="http://www.w3.org/2000/09/xmldsig#"/>'
        declarations = []
        for part, kind in ((origin, 'origin'), (signature, 'xmlsignature+xml')):
            mime = PREFIX + kind if typed else 'application/octet-stream'
            mime = mime.upper() if uppercase else mime
            declarations.append(f'<Override PartName="/{part}" ContentType="{mime}"/>')
        members['[Content_Types].xml'] = members['[Content_Types].xml'].replace(
            b'</Types>', ''.join(declarations).encode('utf-8') + b'</Types>')
        if relationships:
            rel = f'<Relationship Id="signature-origin" Type="{RELATIONSHIP}origin" Target="/{origin}"/>'
            members['_rels/.rels'] = members['_rels/.rels'].replace(
                b'</Relationships>', rel.encode('utf-8') + b'</Relationships>')
            members[folder + '/_rels/origin.sigs.rels'] = (
                f'<Relationships xmlns="{PACKAGE.R}"><Relationship Id="signature" '
                f'Type="{RELATIONSHIP}signature" Target="signature.data"/></Relationships>').encode('utf-8')
        with zipfile.ZipFile(self.source, 'w') as archive:
            for name, data in members.items():
                archive.writestr(name, data)
        return members

    def assert_edit_refused(self):
        original = self.source.read_bytes()
        self.target.write_bytes(b'previous approved delivery')
        with unittest.mock.patch.object(PACKAGE, 'publish_files') as publish:
            with self.assertRaisesRegex(ValueError, 'digital-signatures'):
                PACKAGE.edit_package(self.source, self.target, [{'part': 'word/document.xml',
                    'index': 0, 'expected': 'Original text', 'text': 'Edited text'}])
            publish.assert_not_called()
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(self.target.read_bytes(), b'previous approved delivery')
        self.assertFalse(list(self.root.glob('.revayat-*')))

    def test_noncanonical_signature_parts_are_reported_and_block_edits(self):
        self.fixture()
        self.assertIn('digital-signatures', PACKAGE.inspect_package(self.source)['unsupported'])
        self.assert_edit_refused()

    def test_relationship_signal_alone_cannot_be_hidden_by_generic_content_types(self):
        self.fixture(typed=False)
        self.assert_edit_refused()

    def test_type_signal_alone_is_conservatively_protected(self):
        self.fixture(relationships=False)
        self.assert_edit_refused()

    def test_mime_case_does_not_hide_package_signature_metadata(self):
        self.fixture(relationships=False, uppercase=True)
        self.assert_edit_refused()

    def test_conventional_signature_directory_policy_remains_in_force(self):
        self.fixture(typed=False, relationships=False, folder='_xmlsignatures')
        self.assert_edit_refused()

    def test_readonly_inventory_preserves_metadata_and_document_text(self):
        self.fixture()
        before = self.source.read_bytes()
        report = PACKAGE.inspect_package(self.source)
        self.assertEqual(report['validation'], 'structural-only')
        self.assertEqual(report['render_status'], 'not-performed')
        self.assertIn({'part': 'word/document.xml', 'index': 0, 'text': 'Original text'}, report['text_nodes'])
        self.assertIn('digital-signatures', report['unsupported'])
        self.assertEqual(self.source.read_bytes(), before)

    def test_ordinary_unsigned_package_still_edits_with_exact_other_parts(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Original text')
        document.save(self.source)
        before = self.source.read_bytes()
        with zipfile.ZipFile(self.source) as archive:
            original = {name: archive.read(name) for name in archive.namelist()}
        PACKAGE.edit_package(self.source, self.target, [{'part': 'word/document.xml', 'index': 0,
            'expected': 'Original text', 'text': 'Edited text'}])
        with zipfile.ZipFile(self.target) as archive:
            for name, data in original.items():
                self.assertEqual(archive.read(name), data.replace(b'Original text', b'Edited text')
                                 if name == 'word/document.xml' else data)
        self.assertEqual(self.source.read_bytes(), before)

    def test_arbitrary_signature_word_or_foreign_uri_is_not_opc_signature_metadata(self):
        members = self.fixture(typed=False, relationships=False)
        members['_rels/.rels'] = members['_rels/.rels'].replace(b'</Relationships>',
            b'<Relationship Id="example" Type="https://example.invalid/digital-signature/origin" '
            b'Target="/signing/origin.sigs"/></Relationships>')
        with zipfile.ZipFile(self.source, 'w') as archive:
            for name, data in members.items():
                archive.writestr(name, data)
        self.assertNotIn('digital-signatures', PACKAGE.inspect_package(self.source)['unsupported'])
        PACKAGE.edit_package(self.source, self.target, [{'part': 'word/document.xml', 'index': 0,
            'expected': 'Original text', 'text': 'Edited text'}])
        self.assertTrue(self.target.is_file())


if __name__ == '__main__':
    unittest.main()
