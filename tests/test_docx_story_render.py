"""Render independently expected and edited noncanonical DOCX stories alike."""
import hashlib
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import zipfile

import pymupdf
from docx import Document

from processes import run

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'skills/revayat-scientific/scripts'))
from docx_package import edit_package, inspect_package
from runtime import operation_log


@unittest.skipUnless(os.environ.get('SCIENTIFIC_REQUIRE_DOCX_RENDER')=='1',
                     'native DOCX scientific-render tier')
class DocxStoryRenderTest(unittest.TestCase):
    def test_two_page_renamed_header_footer_matches_independent_expected_render(self):
        office=shutil.which('libreoffice') or shutil.which('soffice')
        self.assertIsNotNone(office,'the required DOCX rendering tier needs LibreOffice Writer')
        with tempfile.TemporaryDirectory(prefix='docx story render ') as directory:
            root=Path(directory).resolve()
            with operation_log('test-docx-story-render',root/'logs') as logger:
                logger.info('creating synthetic two-page source and independent expected bytes')
                canonical=root/'canonical.docx'
                document=Document()
                document.sections[0].header.paragraphs[0].text='HeaderOriginal'
                document.sections[0].footer.paragraphs[0].text='FooterOriginal'
                document.add_paragraph('Scientific body page one: quantity 23.5 C.')
                table=document.add_table(rows=2,cols=2)
                for row,values in zip(table.rows,[('Quantity','Value'),('Count','125')]):
                    for cell,value in zip(row.cells,values):cell.text=value
                document.add_page_break()
                document.add_paragraph('Scientific body page two: unchanged control.')
                document.save(canonical)
                with zipfile.ZipFile(canonical) as archive:
                    members={name:archive.read(name) for name in archive.namelist()}
                aliases={'word/header1.xml':'word/stories/first.xml',
                         'word/footer1.xml':'word/stories/last.xml'}
                for old,new in aliases.items():
                    members[new]=members.pop(old)
                    members['[Content_Types].xml']=members['[Content_Types].xml'].replace(
                        ('/'+old).encode(),('/'+new).encode())
                    members['word/_rels/document.xml.rels']=members['word/_rels/document.xml.rels'].replace(
                        ('Target="'+old[5:]+'"').encode(),('Target="'+new[5:]+'"').encode())
                source,expected,actual=root/'source.docx',root/'expected.docx',root/'actual.docx'
                def write(path,values):
                    with zipfile.ZipFile(path,'w') as archive:
                        for name,data in values.items():archive.writestr(name,data)
                write(source,members)
                before=hashlib.sha256(source.read_bytes()).digest()
                expected_members=dict(members)
                patches=[]
                for part in aliases.values():
                    old='HeaderOriginal' if 'first' in part else 'FooterOriginal'
                    new=old.replace('Original','Updated')
                    expected_members[part]=members[part].replace(old.encode(),new.encode())
                    nodes=[n for n in inspect_package(source)['text_nodes'] if n['part']==part]
                    self.assertEqual(len(nodes),1)
                    patches.append({'part':part,'index':nodes[0]['index'],'expected':old,'text':new})
                write(expected,expected_members)
                edit_package(source,actual,patches)
                with zipfile.ZipFile(actual) as archive:
                    for name,data in expected_members.items():self.assertEqual(archive.read(name),data)
                for index,path in enumerate((expected,actual)):
                    profile=(root/('office-profile-'+str(index))).as_uri()
                    result=run([office,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf',
                                '--outdir',str(root),str(path)],timeout=60)
                    self.assertEqual(result.returncode,0,result.stderr)
                    self.assertTrue(path.with_suffix('.pdf').is_file(),result.stdout)
                with pymupdf.open(expected.with_suffix('.pdf')) as wanted, \
                        pymupdf.open(actual.with_suffix('.pdf')) as got:
                    self.assertEqual(wanted.page_count,2)
                    self.assertEqual(got.page_count,2)
                    for index in range(2):
                        self.assertEqual(wanted[index].get_pixmap().samples,got[index].get_pixmap().samples)
                        extracted=got[index].get_text()
                        self.assertIn('HeaderUpdated',extracted)
                        self.assertIn('FooterUpdated',extracted)
                    evidence=os.environ.get('SCIENTIFIC_EVIDENCE_DIR')
                    if evidence:
                        target=Path(evidence)
                        target.mkdir(parents=True,exist_ok=True)
                        for index,page in enumerate(got):
                            page.get_pixmap().save(target/f'docx-stories-page-{index+1}.png')
                        shutil.copy2(actual.with_suffix('.pdf'),target/'docx-stories.pdf')
                self.assertEqual(hashlib.sha256(source.read_bytes()).digest(),before)
                logger.info('two pages matched exactly; custom stories updated and original unchanged')


if __name__=='__main__':
    unittest.main()
