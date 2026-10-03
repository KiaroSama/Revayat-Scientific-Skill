"""Restricted OOXML features are identified by package metadata, not filenames."""
from pathlib import Path
import json
import sys
import tempfile
import unittest.mock
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills/revayat-scientific/scripts'))
import docx_package as package
from runtime import operation_log
from processes import run
from test_docx import package as base_package

SIG_BASE = 'http://schemas.openxmlformats.org/package/2006/relationships/digital-signature/'
VBA_REL = 'http://schemas.microsoft.com/office/2006/relationships/vbaProject'
FEATURES = (
    ('origin', 'application/vnd.openxmlformats-package.digital-signature-origin', SIG_BASE + 'origin', b''),
    ('signature', 'application/vnd.openxmlformats-package.digital-signature-xmlsignature+xml',
     SIG_BASE + 'signature', b'<Signature xmlns="http://www.w3.org/2000/09/xmldsig#"/>'),
    ('certificate', 'application/vnd.openxmlformats-package.digital-signature-certificate',
     SIG_BASE + 'certificate', b'opaque synthetic certificate, not authenticated'),
    ('project', 'application/vnd.ms-office.vbaProject', VBA_REL, b'opaque synthetic VBA metadata fixture'),
    ('project-signature', 'application/vnd.ms-office.vbaProjectSignature',
     VBA_REL + 'Signature', b'opaque synthetic VBA signature fixture'),
    ('agile-signature', 'application/vnd.ms-office.vbaProjectSignatureAgile',
     VBA_REL + 'SignatureAgile', b'opaque synthetic agile signature fixture'),
    ('agile-canonical', 'application/vnd.ms-office.vbaProjectSignatureAgile',
     'http://schemas.microsoft.com/office/2014/relationships/vbaProjectSignatureAgile', b'synthetic agile metadata'),
    ('v3-signature', 'application/vnd.ms-office.vbaProjectSignatureV3',
     'http://schemas.microsoft.com/office/2020/07/relationships/vbaProjectSignatureV3', b'synthetic V3 metadata'),
)


class DocxFeatureIdentityTest(unittest.TestCase):
    def setUp(self):
        scratch = ROOT / '.scratch'
        scratch.mkdir(exist_ok=True)
        self.scope = tempfile.TemporaryDirectory(prefix='docx feature identity ', dir=scratch)
        self.addCleanup(self.scope.cleanup)
        self.root = Path(self.scope.name).resolve()
        self.source, self.output = self.root / 'source.docx', self.root / 'output.docx'
        context = operation_log('test-docx-feature-identity', self.root / 'logs')
        self.log = context.__enter__()
        self.addCleanup(context.__exit__, None, None, None)
        self.log.info('running test=%s; synthetic metadata is not an actual signature', self._testMethodName)

    def fixture(self, feature, *, typed=True, related=True, mode='Internal', default=False):
        _, content_type, relation, payload = feature
        members = {key: value.encode('utf-8') if isinstance(value, str) else value
                   for key, value in base_package(self.source).items()}
        part = 'opaque/object.dat'
        members[part] = payload
        declared = content_type if typed else 'application/octet-stream'
        entry = (f'<Default Extension="dat" ContentType="{declared}"/>' if default else
                 f'<Override PartName="/{part}" ContentType="{declared}"/>')
        members['[Content_Types].xml'] = members['[Content_Types].xml'].replace(
            b'</Types>', entry.encode() + b'</Types>')
        if related:
            target = part if mode == 'Internal' else 'https://example.invalid/do-not-fetch'
            entry = f'<Relationship Id="restricted" Type="{relation}" Target="{target}" TargetMode="{mode}"/>'
            members['_rels/.rels'] = members['_rels/.rels'].replace(
                b'</Relationships>', entry.encode() + b'</Relationships>')
        self.write(members)
        return members

    def write(self, members):
        with zipfile.ZipFile(self.source, 'w') as archive:
            archive.comment = b'untouched fixture'
            for name, data in members.items():
                archive.writestr(name, data)

    def assert_blocked(self, expected):
        original = self.source.read_bytes()
        self.output.write_bytes(b'previous approved delivery')
        report = package.inspect_package(self.source)
        self.assertIn(expected, report['unsupported'])
        with unittest.mock.patch.object(package, 'publish_files') as publish:
            with self.assertRaises(ValueError):
                package.edit_package(self.source, self.output, [
                    {'part': 'word/document.xml', 'index': 0,
                     'expected': 'Original text', 'text': 'Reviewed replacement'}])
            publish.assert_not_called()
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(self.output.read_bytes(), b'previous approved delivery')
        self.assertFalse(list(self.root.glob('.revayat-*')))

    def test_relocated_signature_parts_are_detected_by_type_and_relationship(self):
        for feature in FEATURES[:3]:
            with self.subTest(part=feature[0]):
                self.fixture(feature)
                self.assert_blocked('digital-signatures')

    def test_signature_content_types_are_sufficient_without_conventional_paths(self):
        for feature in FEATURES[:3]:
            for default in (False, True):
                with self.subTest(part=feature[0], default=default):
                    self.fixture(feature, related=False, default=default)
                    self.assert_blocked('digital-signatures')

    def test_signature_relationships_are_sufficient_with_generic_content_type(self):
        for feature in FEATURES[:3]:
            with self.subTest(part=feature[0]):
                self.fixture(feature, typed=False)
                self.assert_blocked('digital-signatures')

    def test_relocated_vba_parts_are_detected_without_macro_enabled_main(self):
        for feature in FEATURES[3:]:
            for typed, related in ((True, True), (True, False), (False, True)):
                with self.subTest(part=feature[0], typed=typed, related=related):
                    self.fixture(feature, typed=typed, related=related)
                    self.assert_blocked('macros')

    def test_external_restricted_relationship_is_data_not_fetch_permission(self):
        self.fixture(FEATURES[0], typed=False, mode='External')
        self.assert_blocked('digital-signatures')

    def test_override_precedence_does_not_inherit_an_unrelated_restricted_default(self):
        members = self.fixture(FEATURES[0], related=False, default=True)
        members['[Content_Types].xml'] = members['[Content_Types].xml'].replace(
            b'</Types>', b'<Override PartName="/opaque/object.dat" ContentType="application/octet-stream"/></Types>')
        self.write(members)
        self.assertEqual(package.inspect_package(self.source)['unsupported'], [])
        package.edit_package(self.source, self.output, [{'part': 'word/document.xml', 'index': 0,
                              'expected': 'Original text', 'text': 'Changed'}])
        with zipfile.ZipFile(self.output) as result:
            self.assertEqual(result.read('opaque/object.dat'), b'')

    def test_lookalike_custom_types_and_relationships_do_not_grant_edit_restrictions(self):
        feature = ('custom', 'application/example.digital-signature-origin',
                   'https://example.invalid/relationships/digital-signature/origin', b'ordinary custom payload')
        self.fixture(feature)
        report = package.inspect_package(self.source)
        self.assertEqual(report['unsupported'], [])
        package.edit_package(self.source, self.output, [{'part': 'word/document.xml', 'index': 0,
                              'expected': 'Original text', 'text': 'Changed'}])
        self.assertEqual(package.inspect_package(self.output)['text_nodes'][0]['text'], 'Changed')

    def test_existing_canonical_name_and_macro_main_guards_are_retained(self):
        for name, declaration, category in (
            ('_xmlsignatures/legacy.bin', None, 'digital-signatures'),
            ('word/vbaProject.bin', None, 'macros'),
            (None, 'application/vnd.ms-word.document.macroEnabled.main+xml', 'macros')):
            members = {key: value.encode() if isinstance(value, str) else value
                       for key, value in base_package(self.source).items()}
            if name:
                members[name] = b'opaque synthetic legacy metadata'
            if declaration:
                members['[Content_Types].xml'] = members['[Content_Types].xml'].replace(
                    package.MAIN_TYPE.encode(), declaration.encode())
            self.write(members)
            with self.subTest(category=category):
                self.assert_blocked(category)

    def test_unsigned_real_document_stays_editable_without_unrelated_member_changes(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('Original text')
        doc.sections[0].header.paragraphs[0].text = 'Untouched header'
        doc.save(self.source)
        before = self.source.read_bytes()
        with zipfile.ZipFile(self.source) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        package.edit_package(self.source, self.output, [{'part': 'word/document.xml', 'index': 0,
                              'expected': 'Original text', 'text': 'Changed'}])
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(Document(self.output).paragraphs[0].text, 'Changed')
        with zipfile.ZipFile(self.output) as archive:
            for name, value in members.items():
                if name != 'word/document.xml':
                    self.assertEqual(archive.read(name), value)


    def test_mime_type_case_cannot_hide_restricted_features(self):
        for feature in FEATURES:
            modified = (feature[0], feature[1].upper(), feature[2], feature[3])
            with self.subTest(part=feature[0]):
                self.fixture(modified, related=False)
                self.assert_blocked('macros' if feature in FEATURES[3:] else 'digital-signatures')

    def test_complete_relocated_signature_metadata_chain_is_not_dropped(self):
        members = self.fixture(FEATURES[0])
        members['opaque/proof.xml'] = FEATURES[1][3]
        members['opaque/cert.bin'] = FEATURES[2][3]
        extra = ('<Override PartName="/opaque/proof.xml" ContentType="' + FEATURES[1][1] + '"/>'
                 '<Override PartName="/opaque/cert.bin" ContentType="' + FEATURES[2][1] + '"/>')
        members['[Content_Types].xml'] = members['[Content_Types].xml'].replace(
            b'</Types>', extra.encode() + b'</Types>')
        for name, suffix, target in (('object.dat', 'signature', 'proof.xml'),
                                     ('proof.xml', 'certificate', 'cert.bin')):
            members['opaque/_rels/' + name + '.rels'] = (
                '<Relationships xmlns="' + package.R + '"><Relationship Id="r1" Type="'
                + SIG_BASE + suffix + '" Target="' + target + '"/></Relationships>').encode()
        self.write(members)
        self.assert_blocked('digital-signatures')



    def test_public_inspect_reports_restriction_and_edit_preserves_delivery(self):
        self.fixture(FEATURES[0])
        source_bytes = self.source.read_bytes()
        self.output.write_bytes(b'previous approved delivery')
        report, patch = self.root / 'report.json', self.root / 'patch.json'
        command = [sys.executable, str(ROOT / 'skills/revayat-scientific/scripts/revayat-scientific.py')]
        inspected = run([*command, 'docx', 'inspect', str(self.source), '--output', str(report)], timeout=25)
        self.assertEqual(inspected.returncode, 0, inspected.stderr)
        self.assertIn('digital-signatures', json.loads(report.read_text(encoding='utf-8'))['unsupported'])
        patch.write_text(json.dumps([{'part': 'word/document.xml', 'index': 0,
                          'expected': 'Original text', 'text': 'Changed'}]), encoding='utf-8')
        edited = run([*command, 'docx', 'edit', str(self.source), str(self.output), '--patch', str(patch)], timeout=25)
        self.assertNotEqual(edited.returncode, 0)
        self.assertEqual(self.output.read_bytes(), b'previous approved delivery')
        self.assertEqual(self.source.read_bytes(), source_bytes)


if __name__ == '__main__':
    unittest.main()
