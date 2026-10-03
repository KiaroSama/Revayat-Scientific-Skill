"""Address secondary Word stories by declared type, not a generated filename."""
import logging
from pathlib import Path
import sys
import tempfile
import time
import unittest
from unittest import mock
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills/revayat-scientific/scripts'))
import docx_package as package

KINDS = {'header': 'hdr', 'footer': 'ftr', 'footnotes': 'footnotes',
         'endnotes': 'endnotes', 'comments': 'comments'}
TYPE_PREFIX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.'


class DocxStoryPartsTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='docx story audit ')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source, self.output = self.root/'source.docx', self.root/'result.docx'
        self.log = logging.getLogger(self.id())
        handler = logging.StreamHandler()
        formatter = logging.Formatter('%(asctime)s UTC [%(levelname)s] %(message)s')
        formatter.converter = time.gmtime
        handler.setFormatter(formatter)
        self.log.addHandler(handler)
        self.log.setLevel(logging.DEBUG)
        self.addCleanup(self.log.removeHandler, handler)
        self.addCleanup(handler.close)
        self.log.info('running test=%s', self._testMethodName)

    def fixture(self, kind='header', part='word/stories/first.xml', *, generic=False,
                default=False, root=None, extra_type='', story=None):
        mime = 'application/xml' if generic else TYPE_PREFIX + kind + '+xml'
        node = KINDS[kind] if root is None else root
        raw = ('<w:' + node + ' xmlns:w="' + package.W + '"><w:p><w:r>'
               '<w:t>Original</w:t></w:r></w:p></w:' + node + '>').encode()
        decl = ('<Default Extension="'+part.rsplit('.',1)[-1]+'" ContentType="'+mime+'"/>'
                if default else '<Override PartName="/'+part+'" ContentType="'+mime+'"/>')
        types = ('<Types xmlns="'+package.C+'">'
                 '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                 '<Default Extension="xml" ContentType="application/xml"/>'
                 '<Override PartName="/word/document.xml" ContentType="'+package.MAIN_TYPE+'"/>'
                 +decl+extra_type+'</Types>')
        members = {
            '[Content_Types].xml':types.encode(),
            '_rels/.rels': ('<Relationships xmlns="'+package.R+'"><Relationship Id="r1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                'Target="word/document.xml"/></Relationships>').encode(),
            'word/document.xml': ('<w:document xmlns:w="'+package.W+'"><w:body><w:p><w:r>'
                '<w:t>Body</w:t></w:r></w:p></w:body></w:document>').encode(),
            'word/_rels/document.xml.rels': ('<Relationships xmlns="'+package.R+'"><Relationship Id="s1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'+kind+'" '
                'Target="/'+part+'"/></Relationships>').encode(),
            part: raw if story is None else story,
        }
        self.write(members)
        return members

    def write(self, members):
        with zipfile.ZipFile(self.source, 'w') as archive:
            archive.comment = b'untouched package comment'
            # Reverse order to prove discovery does not depend on ZIP member order.
            for name, data in reversed(list(members.items())):
                archive.writestr(name, data)

    def test_all_five_story_kinds_use_declared_names(self):
        for kind in KINDS:
            for part in ('word/stories/named.xml', 'other/named.part', 'named'):
                with self.subTest(kind=kind,part=part):
                    self.fixture(kind,part)
                    nodes=package.inspect_package(self.source)['text_nodes']
                    self.assertIn({'part':part,'index':0,'text':'Original'},nodes)

    def test_inspected_address_is_editable_without_changing_other_bytes(self):
        for kind in KINDS:
            with self.subTest(kind=kind):
                part='shared/'+kind+'.story'
                before=self.fixture(kind,part)
                source_bytes=self.source.read_bytes()
                patch=[{'part':part,'index':0,'expected':'Original','text':' متن & <x> '}]
                package.edit_package(self.source,self.output,patch)
                with zipfile.ZipFile(self.output) as archive:
                    self.assertEqual(set(archive.namelist()),set(before))
                    self.assertEqual(archive.comment,b'untouched package comment')
                    for name,data in before.items():
                        expected=(data.replace(b'<w:t>Original</w:t>',
                            '<w:t xml:space="preserve"> متن &amp; &lt;x&gt; </w:t>'.encode())
                            if name==part else data)
                        self.assertEqual(archive.read(name),expected)
                self.assertEqual(self.source.read_bytes(),source_bytes)

    def test_default_story_type_and_override_precedence(self):
        part='word/custom.hdr'
        self.fixture(part=part,default=True)
        self.assertIn(part,{n['part'] for n in package.inspect_package(self.source)['text_nodes']})
        self.fixture(part=part,default=True,
            extra_type='<Override PartName="/'+part+'" ContentType="application/xml"/>')
        self.assertNotIn(part,{n['part'] for n in package.inspect_package(self.source)['text_nodes']})

    def test_generic_custom_xml_is_not_authorized_by_a_text_node(self):
        part='customXml/item.xml'
        self.fixture(part=part,generic=True)
        self.assertNotIn(part,{n['part'] for n in package.inspect_package(self.source)['text_nodes']})
        with self.assertRaises(ValueError):
            package.edit_package(self.source,self.output,[{'part':part,'index':0,
                'expected':'Original','text':'Changed'}])
        self.assertFalse(self.output.exists())

    def test_legacy_canonical_generic_story_remains_supported(self):
        part='word/header1.xml'
        self.fixture(part=part,generic=True)
        package.edit_package(self.source,self.output,[{'part':part,'index':0,
            'expected':'Original','text':'Changed'}])
        self.assertIn({'part':part,'index':0,'text':'Changed'},
            package.inspect_package(self.output)['text_nodes'])

    def test_wrong_declared_story_root_fails_before_publication(self):
        for root in ('document','ftr','custom'):
            with self.subTest(root=root):
                self.fixture(root=root)
                self.output.write_bytes(b'previous delivery')
                with self.assertRaises(ValueError):
                    package.edit_package(self.source,self.output,[{'part':'word/stories/first.xml',
                        'index':0,'expected':'Original','text':'Changed'}])
                self.assertEqual(self.output.read_bytes(),b'previous delivery')

    def test_duplicate_declarations_are_not_silently_selected(self):
        self.fixture(extra_type='<Override PartName="/word/stories/first.xml" ContentType="application/xml"/>')
        with self.assertRaises(ValueError): package.read_package(self.source)

    def test_typed_nonxml_suffix_cannot_bypass_xml_limits(self):
        samples=(b'<!DOCTYPE hdr><w:hdr xmlns:w="'+package.W.encode()+b'"/>',
                 b'\xff', b'<broken', b' '*(package.MAX_XML+1))
        for data in samples:
            with self.subTest(size=len(data)):
                self.fixture(part='word/first.story',story=data)
                with self.assertRaises((ValueError,package.ET.ParseError)):
                    package.read_package(self.source)

    def test_late_stale_edit_keeps_source_and_destination(self):
        part='word/custom.xml'
        self.fixture(part=part)
        before=self.source.read_bytes()
        self.output.write_bytes(b'previous delivery')
        with self.assertRaisesRegex(ValueError,'stale'):
            package.edit_package(self.source,self.output,[
                {'part':'word/document.xml','index':0,'expected':'Body','text':'New body'},
                {'part':part,'index':0,'expected':'Stale','text':'New header'}])
        self.assertEqual(self.output.read_bytes(),b'previous delivery')
        self.assertEqual(self.source.read_bytes(),before)
        self.assertFalse(list(self.root.glob('.revayat-*')))

    def test_duplicate_addresses_and_publication_failure_are_recoverable(self):
        part='word/custom.xml'
        self.fixture(part=part)
        patch={'part':part,'index':0,'expected':'Original','text':'Changed'}
        self.output.write_bytes(b'previous delivery')
        with self.assertRaises(ValueError): package.edit_package(self.source,self.output,[patch,patch])
        with mock.patch('publication.os.replace',side_effect=OSError('injected failure')):
            with self.assertRaises(OSError): package.edit_package(self.source,self.output,[patch])
        self.assertEqual(self.output.read_bytes(),b'previous delivery')
        self.assertFalse(list(self.root.glob('.revayat-*')))

    def test_python_docx_reads_and_edits_a_renamed_real_header(self):
        from docx import Document
        document=Document()
        document.add_paragraph('Body')
        document.sections[0].header.paragraphs[0].text='Original'
        document.save(self.source)
        with zipfile.ZipFile(self.source) as archive:
            members={n:archive.read(n) for n in archive.namelist()}
        part='word/stories/first.xml'
        members[part]=members.pop('word/header1.xml')
        members['[Content_Types].xml']=members['[Content_Types].xml'].replace(
            b'/word/header1.xml',('/'+part).encode())
        members['word/_rels/document.xml.rels']=members['word/_rels/document.xml.rels'].replace(
            b'Target="header1.xml"',b'Target="stories/first.xml"')
        self.write(members)
        self.assertEqual(Document(self.source).sections[0].header.paragraphs[0].text,'Original')
        package.edit_package(self.source,self.output,[{'part':part,'index':0,
            'expected':'Original','text':'Changed'}])
        self.assertEqual(Document(self.output).sections[0].header.paragraphs[0].text,'Changed')


if __name__ == '__main__':
    unittest.main()
