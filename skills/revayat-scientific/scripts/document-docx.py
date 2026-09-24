#!/usr/bin/env python3
"""Inspect, create, edit and structurally validate DOCX without external skills."""
import argparse
import json
import math
from pathlib import Path
import re
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from docx_package import edit_package, inspect_package, read_package, require_editable
from publication import publish_files, validate_destination
from runtime import operation_log


def _unique(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError('duplicate JSON key')
        result[key] = value
    return result


def load_json(path):
    if path.stat().st_size > 8 * 1024 * 1024:
        raise ValueError('JSON input exceeds 8 MiB limit')
    return json.loads(path.read_text(encoding='utf-8'), object_pairs_hook=_unique)


def _object(value, required, optional=()):
    if not isinstance(value, dict) or not set(required) <= set(value) or set(value) - set(required) - set(optional):
        raise ValueError('missing or unknown content JSON fields')
    return value


def _text(value):
    if (not isinstance(value, str) or len(value) > 100000
            or any(ord(char) < 32 and char not in '\t\n' for char in value)
            or any(0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in value)):
        raise ValueError('invalid document text')
    return value


def _number(value, minimum, maximum):
    if type(value) not in (int, float) or not math.isfinite(value) or not minimum <= value <= maximum:
        raise ValueError('numeric content value is outside supported limits')
    return value


def _boolean(value):
    if type(value) is not bool:
        raise ValueError('direction and style flags must be JSON booleans')
    return value


def _language(value):
    if not isinstance(value, str) or not re.fullmatch(r'[A-Za-z]{2,8}(?:-[A-Za-z0-9]{1,8})*', value):
        raise ValueError('invalid language tag')
    return value


def _paragraph(parent, spec, language, rtl):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    if isinstance(spec, str):
        spec = {'text': spec}
    _object(spec, (), ('type', 'text', 'runs', 'style', 'language', 'rtl'))
    if ('text' in spec) == ('runs' in spec):
        raise ValueError('paragraph needs exactly text or runs')
    language = _language(spec.get('language', language))
    rtl = _boolean(spec.get('rtl', rtl))
    style = _text(spec['style']) if 'style' in spec else None
    paragraph = parent.add_paragraph(style=style)
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1' if rtl else '0')
    paragraph._p.get_or_add_pPr().append(bidi)
    runs = spec.get('runs', [{'text': spec.get('text', '')}])
    if not isinstance(runs, list) or len(runs) > 10000:
        raise ValueError('runs must be an array of at most 10000 entries')
    for item in runs:
        _object(item, ('text',), ('bold', 'italic', 'font', 'size_pt', 'language', 'rtl'))
        run = paragraph.add_run(_text(item['text']))
        run.font.rtl = _boolean(item.get('rtl', rtl))
        for flag in ('bold', 'italic'):
            if flag in item:
                setattr(run, flag, _boolean(item[flag]))
        if 'font' in item:
            run.font.name = _text(item['font'])
            fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn('w:cs'), item['font'])
        if 'size_pt' in item:
            size = _number(item['size_pt'], 1, 200)
            run.font.size = Pt(size)
            complex_size = OxmlElement('w:szCs')
            complex_size.set(qn('w:val'), str(round(size * 2)))
            run._r.get_or_add_rPr().append(complex_size)
        lang = OxmlElement('w:lang')
        selected = _language(item.get('language', language))
        lang.set(qn('w:val'), selected)
        lang.set(qn('w:bidi'), selected)
        run._r.get_or_add_rPr().append(lang)
    return paragraph


def _image(root, value):
    if not isinstance(value, str) or not value or Path(value).is_absolute() or '..' in Path(value).parts:
        raise ValueError('image must name an approved file relative to the content JSON')
    image = root / value
    if any(part.is_symlink() or (hasattr(part, 'is_junction') and part.is_junction())
           for part in (image, *image.parents)):
        raise ValueError('image path must not contain links or junctions')
    if not image.resolve().is_relative_to(root.resolve()) or not image.is_file():
        raise ValueError('image is missing or escapes the content directory')
    if image.stat().st_size > 64 * 1024 * 1024:
        raise ValueError('image exceeds 64 MiB limit')
    from PIL import Image
    try:
        with Image.open(image) as asset:
            if (asset.format not in ('PNG', 'JPEG') or getattr(asset, 'n_frames', 1) != 1
                    or asset.width * asset.height > 50_000_000):
                raise ValueError('create supports single-frame PNG/JPEG up to 50 million pixels')
            if asset.getexif().get(274, 1) != 1:
                raise ValueError('normalize a copy of the image orientation before embedding')
        with Image.open(image) as asset:
            asset.verify()
    except Image.DecompressionBombError as error:
        raise ValueError('image exceeds decompression limits') from error
    return image


def create_document(content_path, destination):
    # python-docx is used only for new documents, never to round-trip source files.
    from docx import Document
    from docx.enum.section import WD_SECTION_START, WD_ORIENT
    from docx.enum.table import WD_TABLE_DIRECTION
    from docx.shared import Mm

    content = _object(load_json(content_path), ('sections',), ('language', 'rtl'))
    sections = content['sections']
    if not isinstance(sections, list) or not sections or len(sections) > 128:
        raise ValueError('sections must contain 1 to 128 entries')
    language = _language(content.get('language', 'fa-IR'))
    rtl = _boolean(content.get('rtl', True))
    destination = validate_destination(destination, (content_path,))
    document = Document()
    protected = [content_path]
    for index, spec in enumerate(sections):
        _object(spec, ('width_mm', 'height_mm', 'margins_mm', 'blocks'), ('header', 'footer'))
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        width, height = (_number(spec[key], 25, 2000) for key in ('width_mm', 'height_mm'))
        section.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Mm(width), Mm(height)
        margins = _object(spec['margins_mm'], ('top', 'bottom', 'left', 'right'))
        for side, value in margins.items():
            setattr(section, side + '_margin', Mm(_number(value, 0, 1000)))
        if margins['left'] + margins['right'] >= width or margins['top'] + margins['bottom'] >= height:
            raise ValueError('section margins leave no usable page area')
        for kind in ('header', 'footer'):
            story = getattr(section, kind)
            story.is_linked_to_previous = False
            if kind in spec:
                paragraph = _paragraph(story, spec[kind], language, rtl)
                # New headers/footers contain one mandatory empty paragraph.
                empty = story.paragraphs[0]._p
                if empty is not paragraph._p:
                    empty.getparent().remove(empty)
        if not isinstance(spec['blocks'], list) or len(spec['blocks']) > 10000:
            raise ValueError('blocks must be an array of at most 10000 entries')
        for block in spec['blocks']:
            if not isinstance(block, dict):
                raise ValueError('each block must be an object with a type')
            kind = block.get('type')
            if kind == 'paragraph':
                _paragraph(document, block, language, rtl)
            elif kind == 'table':
                _object(block, ('type', 'rows'), ('language', 'rtl'))
                rows = block['rows']
                if (not isinstance(rows, list) or not rows or len(rows) > 1000
                        or not isinstance(rows[0], list) or not 1 <= len(rows[0]) <= 100
                        or any(not isinstance(row, list) or len(row) != len(rows[0]) for row in rows)):
                    raise ValueError('table must have 1..1000 rectangular rows and 1..100 columns')
                table = document.add_table(rows=len(rows), cols=len(rows[0]))
                table_rtl = _boolean(block.get('rtl', rtl))
                table.table_direction = WD_TABLE_DIRECTION.RTL if table_rtl else WD_TABLE_DIRECTION.LTR
                for cells, values in zip(table.rows, rows):
                    for cell, value in zip(cells.cells, values):
                        empty = cell.paragraphs[0]._p
                        _paragraph(cell, value, _language(block.get('language', language)), table_rtl)
                        empty.getparent().remove(empty)
            elif kind == 'image':
                _object(block, ('type', 'path', 'width_mm'))
                asset = _image(content_path.resolve().parent, block['path'])
                protected.append(asset)
                image_width = _number(block['width_mm'], 1, width - margins['left'] - margins['right'])
                document.add_picture(str(asset), width=Mm(image_width))
            else:
                raise ValueError('unsupported block type; use paragraph, table or image')
    validate_destination(destination, protected)
    with tempfile.TemporaryDirectory(prefix='.revayat-docx-', dir=destination.parent) as directory:
        staged = Path(directory) / 'created.docx'
        document.save(staged)
        read_package(staged)
        publish_files([(staged, destination)], protected_sources=protected)


def write_report(report, destination, source):
    destination = validate_destination(destination, (source,))
    with tempfile.TemporaryDirectory(prefix='.revayat-docx-', dir=destination.parent) as directory:
        stage = Path(directory) / 'report.json'
        stage.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        publish_files([(stage, destination)], protected_sources=(source,))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    inspect = commands.add_parser('inspect', help='write structure and addressable text JSON')
    inspect.add_argument('input', type=Path)
    inspect.add_argument('--output', type=Path, required=True)
    edit = commands.add_parser('edit', help='apply exact expected-text patches to a new file')
    edit.add_argument('input', type=Path)
    edit.add_argument('output', type=Path)
    edit.add_argument('--patch', type=Path, required=True)
    create = commands.add_parser('create', help='create a new DOCX from content JSON')
    create.add_argument('input', type=Path)
    create.add_argument('output', type=Path)
    validate = commands.add_parser('validate', help='validate supported package structure, not rendering')
    validate.add_argument('input', type=Path)
    args = parser.parse_args(argv)
    try:
        with operation_log('document-docx', Path(__file__).resolve().parent / 'logs') as logger:
            logger.debug('operation=%s', args.command)
            if args.command == 'inspect':
                report = inspect_package(args.input)
                write_report(report, args.output, args.input)
                if report['unsupported']:
                    logger.warning('unsupported_feature_count=%d', len(report['unsupported']))
                logger.info('inspected text_nodes=%d media=%d', len(report['text_nodes']), len(report['media']))
            elif args.command == 'edit':
                validate_destination(args.output, (args.input, args.patch))
                patches = load_json(args.patch)
                edit_package(args.input, args.output, patches, protected_sources=(args.patch,))
                logger.info('edited patch_count=%d', len(patches))
            elif args.command == 'create':
                create_document(args.input, args.output)
                logger.info('created document')
            else:
                report = inspect_package(args.input)
                require_editable(report['unsupported'])
                logger.info('validated structure; rendering not performed')
            print('DOCX operation completed; layout/render verification remains separate.')
            return 0
    except ImportError:
        print('document-docx: creation requires the installed skill requirements (python-docx and Pillow).', file=sys.stderr)
        return 2
    except (ValueError, UnicodeError, ET.ParseError) as error:
        # Known validation messages contain no document text or external values.
        message = str(error) if type(error) is ValueError else 'invalid document or JSON encoding/structure'
        print('document-docx: ' + message, file=sys.stderr)
        return 2
    except (OSError, RuntimeError, zipfile.BadZipFile, KeyError) as error:
        if isinstance(error, RuntimeError) and str(error).startswith('publication rollback needs recovery: '):
            print('document-docx: ' + str(error), file=sys.stderr)
            return 1
        print('document-docx: operation failed (' + type(error).__name__ + '); original inputs preserved.', file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
